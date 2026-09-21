import sys, os, pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

src_path = r'd:\folder\tools\other\canva\Slide_baocao.pptx'
dst_path = r'd:\folder\tools\other\canva\Slide_baocao_DauCau_v2.pptx'

prs = pptx.Presentation(src_path)

COLOR_NAVY = (30, 58, 138)         # #1E3A8A - IUH Deep Navy
COLOR_BLUE = (15, 76, 129)         # #0F4C81 - Classic IUH Blue
COLOR_TEXT_MAIN = (30, 41, 59)     # #1E293B - Slate Dark
COLOR_TEXT_MUTED = (100, 116, 139) # #64748B - Slate Muted
COLOR_CARD_BG = (255, 255, 255)    # #FFFFFF
COLOR_CARD_BORDER = (218, 226, 237)# Soft card border
COLOR_ACCENT = (37, 99, 235)       # Royal Blue for emphasis
COLOR_GREEN = (22, 101, 52)        # Forest Green for correct examples
COLOR_RED = (185, 28, 28)          # Red for wrong examples

def set_font(run, name='Canva Sans', size_pt=18, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)

def clear_and_set_text(text_frame, text, font_name='Canva Sans', size_pt=18, bold=False, italic=False, color_rgb=(0,0,0), align=PP_ALIGN.LEFT):
    text_frame.text = ""
    p = text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, font_name, size_pt, bold, italic, color_rgb)
    return p

def clean_slide_content_shapes(slide, keep_names=['Freeform 2']):
    for sh in list(slide.shapes):
        if sh.name not in keep_names:
            sp = sh._element
            sp.getparent().remove(sp)

def setup_header(slide, title_text):
    tb_title = slide.shapes.add_textbox(Inches(1.3), Inches(1.85), Inches(15.3), Inches(0.80))
    clear_and_set_text(tb_title.text_frame, title_text, 'Canva Sans Bold', 26, True, False, COLOR_NAVY, PP_ALIGN.LEFT)

def create_card(slide, left, top, width, height, bg_rgb=COLOR_CARD_BG, border_rgb=COLOR_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(*bg_rgb)
    card.line.color.rgb = RGBColor(*border_rgb)
    card.line.width = Pt(1.5)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.40)
    tf.margin_right = Inches(0.40)
    tf.margin_top = Inches(0.35)
    tf.margin_bottom = Inches(0.35)
    return card

def populate_card_content(card, title, items, title_size=20, item_size=17, line_spacing=5):
    tf = card.text_frame
    tf.text = ""
    
    if title:
        p0 = tf.paragraphs[0]
        p0.alignment = PP_ALIGN.LEFT
        p0.space_after = Pt(10)
        r0 = p0.add_run()
        r0.text = title
        set_font(r0, 'Canva Sans Bold', title_size, True, False, COLOR_NAVY)
    
    for idx, item in enumerate(items):
        if not title and idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(line_spacing)
        p.space_after = Pt(3)
        
        if item.startswith("• "):
            r_bullet = p.add_run()
            r_bullet.text = "• "
            set_font(r_bullet, 'Canva Sans Bold', item_size, True, False, COLOR_BLUE)
            content = item[2:]
            if ":" in content:
                heading, rest = content.split(":", 1)
                r_h = p.add_run()
                r_h.text = heading + ":"
                set_font(r_h, 'Canva Sans Bold', item_size, True, False, COLOR_NAVY)
                r_t = p.add_run()
                r_t.text = rest
                set_font(r_t, 'Canva Sans', item_size, False, False, COLOR_TEXT_MAIN)
            else:
                r_t = p.add_run()
                r_t.text = content
                set_font(r_t, 'Canva Sans', item_size, False, False, COLOR_TEXT_MAIN)
        elif item.startswith("  - ") or item.startswith("-> ") or item.startswith("- "):
            p.space_before = Pt(2)
            r_sub = p.add_run()
            r_sub.text = "    " + item.strip()
            if "Đúng" in item or "✅" in item:
                set_font(r_sub, 'Canva Sans Bold', item_size, True, False, COLOR_GREEN)
            elif "Sai" in item or "❌" in item:
                set_font(r_sub, 'Canva Sans', item_size, False, False, COLOR_RED)
            else:
                set_font(r_sub, 'Canva Sans', item_size, False, False, COLOR_TEXT_MAIN)
        elif item.startswith("Ví dụ "):
            if ":" in item:
                h_ex, text_ex = item.split(":", 1)
                r_ex = p.add_run()
                r_ex.text = h_ex + ":"
                set_font(r_ex, 'Canva Sans Bold', item_size, True, False, COLOR_NAVY)
                r_body = p.add_run()
                r_body.text = text_ex
                set_font(r_body, 'Canva Sans', item_size, False, False, COLOR_TEXT_MAIN)
            else:
                r_ex = p.add_run()
                r_ex.text = item
                set_font(r_ex, 'Canva Sans Bold', item_size, True, False, COLOR_NAVY)
        else:
            r_t = p.add_run()
            r_t.text = item
            set_font(r_t, 'Canva Sans', item_size, False, False, COLOR_TEXT_MAIN)

print("=== REBUILDING 14-SLIDE LARGE & CLEAR PRESENTATION ===")

# Relink slide 13 (index 12) background to clean IUH background (image3.jpeg)
img3_part = prs.slides[1].part.rels['rId2'].target_part
new_rId = prs.slides[12].part.relate_to(img3_part, pptx.opc.constants.RELATIONSHIP_TYPE.IMAGE)
for blip in prs.slides[12].shapes[0]._element.xpath('.//a:blip'):
    blip.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'] = new_rId

# Add slide 14 with image3.jpeg background if prs has only 13 slides
if len(prs.slides) == 13:
    s14_new = prs.slides.add_slide(prs.slide_layouts[6])
    pic14 = s14_new.shapes.add_picture(r'd:\folder\tools\other\canva\extracted_media\image3.jpeg', 0, 0, Inches(20), Inches(11.25))
    pic14.name = 'Freeform 2'

# Ensure exactly 14 slides
while len(prs.slides) > 14:
    rId = prs.slides._sldIdLst[len(prs.slides)-1].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[len(prs.slides)-1]

print(f"Working on exactly {len(prs.slides)} slides...")

# =============================================================
# SLIDE 1: Bìa bài thuyết trình (Tự nhiên, không AI)
# =============================================================
print("Slide 1: Bìa bài thuyết trình...")
s1 = prs.slides[0]
clean_slide_content_shapes(s1)

tb_c_tag = s1.shapes.add_textbox(Inches(1.5), Inches(1.75), Inches(15.288), Inches(0.5))
clear_and_set_text(tb_c_tag.text_frame, "TIỂU LUẬN HỌC PHẦN: TIẾNG VIỆT THỰC HÀNH", 'Canva Sans Bold', 17, True, False, COLOR_BLUE, PP_ALIGN.CENTER)

tb_c_title = s1.shapes.add_textbox(Inches(1.2), Inches(2.35), Inches(15.888), Inches(1.8))
tf_title = tb_c_title.text_frame
tf_title.text = ""
p1 = tf_title.paragraphs[0]
p1.alignment = PP_ALIGN.CENTER
r1 = p1.add_run()
r1.text = "TÌM HIỂU VỀ HỆ THỐNG DẤU CÂU\nTRONG TIẾNG VIỆT"
set_font(r1, 'Canva Sans Bold', 40, True, False, COLOR_NAVY)

card_cover = create_card(s1, Inches(2.2), Inches(4.35), Inches(13.888), Inches(3.95), (255,255,255), (203,213,225))
tf_cc = card_cover.text_frame
tf_cc.text = ""

p_c1 = tf_cc.paragraphs[0]
p_c1.alignment = PP_ALIGN.CENTER
r_c1 = p_c1.add_run()
r_c1.text = "ĐỀ TÀI: TÌM HIỂU VAI TRÒ, TÁC DỤNG, VỊ TRÍ VÀ VÍ DỤ CỦA CÁC DẤU CÂU THÔNG DỤNG"
set_font(r_c1, 'Canva Sans Bold', 19, True, False, COLOR_BLUE)

p_c2 = tf_cc.add_paragraph()
p_c2.alignment = PP_ALIGN.CENTER
p_c2.space_before = Pt(8)
r_c2 = p_c2.add_run()
r_c2.text = "Lớp: DHCNTP21A — Nhóm thực hiện: Nhóm 5"
set_font(r_c2, 'Canva Sans Bold', 17, True, False, COLOR_NAVY)

p_c3 = tf_cc.add_paragraph()
p_c3.alignment = PP_ALIGN.CENTER
p_c3.space_before = Pt(12)
r_c3 = p_c3.add_run()
r_c3.text = "Nội dung tìm hiểu gồm 5 dấu câu cơ bản:\n• Phần 1: Dấu Chấm (.) — Dấu Phẩy (,) — Dấu Hỏi (?)\n• Phần 2: Dấu Chấm Than (!) — Dấu Hai Chấm (:)"
set_font(r_c3, 'Canva Sans', 16, False, False, COLOR_TEXT_MAIN)

p_c4 = tf_cc.add_paragraph()
p_c4.alignment = PP_ALIGN.CENTER
p_c4.space_before = Pt(12)
r_c4 = p_c4.add_run()
r_c4.text = "Giảng viên hướng dẫn: Bộ môn Tiếng Việt thực hành"
set_font(r_c4, 'Canva Sans Italic', 15, False, True, COLOR_TEXT_MUTED)

# =============================================================
# SLIDE 2: Danh sách thành viên Nhóm 5 (BỎ SĐT)
# =============================================================
print("Slide 2: Danh sách thành viên Nhóm 5...")
s2 = prs.slides[1]
clean_slide_content_shapes(s2)
setup_header(s2, "Danh sách thành viên Nhóm 5 — Lớp: DHCNTP21A")

sh2_tbl = s2.shapes.add_table(9, 4, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
t2 = sh2_tbl.table
grid2 = sh2_tbl._element.xpath('.//a:tblGrid')[0]
col_w2 = [1097280, 4389120, 3017520, 5486400]
for idx, w in enumerate(col_w2):
    grid2.xpath('.//a:gridCol')[idx].attrib['w'] = str(w)
    t2.columns[idx].width = w

headers2 = ["STT", "Họ và tên", "MSSV", "Nhiệm vụ phân công trong bài"]
members2 = [
    ["1", "Trịnh Quang Hiên", "23632591", "Nhóm trưởng | Thuyết trình & Duyệt nội dung toàn bài"],
    ["2", "Nguyễn Văn Nhật", "24725801", "Nội dung 2: Tìm hiểu Dấu hai chấm (:)"],
    ["3", "Nguyễn Văn Minh", "25688761", "Nội dung 1: Tìm hiểu Dấu phẩy (,)"],
    ["4", "Nguyễn Ngọc Lợi", "25708311", "Nội dung 1: Tìm hiểu Dấu chấm (.)"],
    ["5", "Võ Khương Duy", "23646941", "Làm Slide Nội dung 2 & Slide kết thúc"],
    ["6", "Lê Đình Công", "24727921", "Nội dung 2: Tìm hiểu Dấu chấm than (!)"],
    ["7", "Lê Thành Long", "23630851", "Làm Slide Giới thiệu, Thành viên & Nội dung 1"],
    ["8", "Trương Bảo Huy", "25639981", "Nội dung 1: Tìm hiểu Dấu hỏi (?)"],
]

for c_idx, h in enumerate(headers2):
    cell = t2.cell(0, c_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(*COLOR_NAVY)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    clear_and_set_text(cell.text_frame, h, 'Canva Sans Bold', 18, True, False, (255, 255, 255), PP_ALIGN.CENTER)

for r_idx, row_data in enumerate(members2):
    bg = RGBColor(255, 255, 255) if r_idx % 2 == 0 else RGBColor(241, 245, 249)
    for c_idx, val in enumerate(row_data):
        cell = t2.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        align = PP_ALIGN.LEFT if c_idx in [1, 3] else PP_ALIGN.CENTER
        is_user = "Lê Thành Long" in val
        is_leader = "Nhóm trưởng" in val
        bold = True if (c_idx == 1 or is_user or is_leader) else False
        color = COLOR_NAVY if (is_user or is_leader) else COLOR_TEXT_MAIN
        clear_and_set_text(cell.text_frame, val, 'Canva Sans', 16, bold, False, color, align)

# =============================================================
# SLIDE 3: I. Đặt vấn đề: Tầm quan trọng của dấu câu tiếng Việt
# =============================================================
print("Slide 3: Đặt vấn đề - Tầm quan trọng của dấu câu...")
s3 = prs.slides[2]
clean_slide_content_shapes(s3)
setup_header(s3, "I. Đặt vấn đề: Tầm quan trọng của dấu câu tiếng Việt")

card3 = create_card(s3, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s3_items = [
    "• Dấu câu là phương tiện diễn đạt trực quan trong văn bản viết:",
    "Dấu câu giúp người viết phân định rõ ràng các vế câu, tạo quãng ngắt nghỉ nhịp nhàng và hướng dẫn người đọc tiếp thu đúng mạch ý.",
    "• Tầm quan trọng của việc dùng đúng dấu câu:",
    "Dùng đúng dấu câu giúp câu văn trong sáng, mạch lạc, thể hiện sự chỉn chu và tránh gây hiểu lầm tai hại trong giao tiếp.",
    "• Ví dụ thực tế về việc đặt sai dấu câu làm thay đổi hẳn ý nghĩa:",
    "  - Khi viết sai: \"Bố hút thuốc, lá rất có hại cho sức khỏe.\" (Nghĩa bị hiểu nhầm sang lá cây)",
    "  - Khi viết đúng: \"Bố hút thuốc lá, rất có hại cho sức khỏe.\" (Rõ nghĩa, chính xác)",
    "• Mục tiêu của bài thuyết trình:",
    "Giúp các bạn sinh viên nắm vững vị trí, công dụng và quy tắc viết chuẩn của từng dấu câu, từ đó áp dụng hiệu quả vào bài tập và giao tiếp hàng ngày."
]
populate_card_content(card3, "Tầm quan trọng của hệ thống dấu câu tiếng Việt", s3_items, title_size=21, item_size=18, line_spacing=6)

# =============================================================
# SLIDE 4: Bố cục nội dung bài thuyết trình (Không để II. lạc lõng)
# =============================================================
print("Slide 4: Bố cục nội dung bài thuyết trình...")
s4 = prs.slides[3]
clean_slide_content_shapes(s4)
setup_header(s4, "Bố cục nội dung bài thuyết trình")

card4 = create_card(s4, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s4_items = [
    "• I. ĐẶT VẤN ĐỀ:",
    "Tầm quan trọng của hệ thống dấu câu và các hiểu lầm tai hại khi đặt sai dấu trong câu văn.",
    "• PHẦN 1 — TÌM HIỂU CHI TIẾT 3 DẤU CÂU CƠ BẢN (Nội dung báo cáo hôm nay):",
    "  - II. Dấu Phẩy (,): Vị trí trong câu — Quy tắc khoảng trắng — Công dụng — 8 Ví dụ thực tế.",
    "  - III. Dấu Chấm (.): Vị trí kết thúc câu kể — Quy ước viết tắt, số tiền, ngày tháng — 8 Ví dụ thực tế.",
    "  - IV. Dấu Hỏi (?): Vị trí nghi vấn — Quy tắc viết — Ngữ điệu & câu hỏi tu từ — 8 Ví dụ thực tế.",
    "• PHẦN 2 — TÌM HIỂU 2 DẤU CÂU TIẾP THEO (Phần tiếp nối của nhóm):",
    "  - V. Dấu Chấm Than (!): Bộc lộ cảm xúc mãnh liệt, câu cầu khiến, ra lệnh hoặc lời hô gọi.",
    "  - VI. Dấu Hai Chấm (:): Báo hiệu giải thích, thuyết minh, dẫn lời nói trực tiếp hoặc mở đầu liệt kê.",
    "• VII. TỔNG KẾT & BÀI HỌC VẬN DỤNG:",
    "Đúc kết các quy tắc soạn thảo văn bản chuẩn và liên hệ thực tế trong học tập."
]
populate_card_content(card4, "Hệ thống các phần được trình bày trong bài báo cáo", s4_items, title_size=21, item_size=16, line_spacing=4)

# =============================================================
# SLIDE 5: II. Dấu phẩy: Vị trí trong câu (Card + Ảnh 3D)
# =============================================================
print("Slide 5: Dấu phẩy - Vị trí trong câu...")
s5 = prs.slides[4]
clean_slide_content_shapes(s5)
setup_header(s5, "II. Dấu phẩy: Vị trí xuất hiện trong câu")

card5 = create_card(s5, Inches(1.3), Inches(2.80), Inches(9.6), Inches(5.75))
s5_items = [
    "• Đặt ngay sau từ hoặc cụm từ cần ngắt nhịp:",
    "Dấu phẩy luôn bám sát từ ngữ cần tạo quãng nghỉ trong câu.",
    "• Đặt sau trạng ngữ ở đầu câu:",
    "Ngăn cách trạng ngữ chỉ thời gian, địa điểm, nguyên nhân với cụm Chủ - Vị chính.",
    "• Đặt giữa câu để ngăn cách chuỗi liệt kê:",
    "Tách các từ, cụm từ có cùng chức vụ ngữ pháp đứng liền nhau.",
    "• Đặt giữa các vế trong câu ghép:",
    "Làm ranh giới phân định các vế câu độc lập về nghĩa.",
    "• Đặt bao quanh thành phần chú thích hoặc lời gọi đáp:",
    "Cô lập phần giải thích hoặc hô ngữ xen vào giữa câu."
]
populate_card_content(card5, "Các vị trí thường gặp của Dấu Phẩy", s5_items, title_size=21, item_size=18, line_spacing=6)

card5_img = create_card(s5, Inches(11.3), Inches(2.80), Inches(5.3), Inches(5.75))
s5.shapes.add_picture(r'd:\folder\tools\other\canva\assets\dau_phay.jpg', Inches(11.45), Inches(2.95), Inches(5.0), Inches(5.0))
tb5_lbl = s5.shapes.add_textbox(Inches(11.3), Inches(8.0), Inches(5.3), Inches(0.45))
clear_and_set_text(tb5_lbl.text_frame, "DẤU PHẨY (,)", 'Canva Sans Bold', 17, True, False, COLOR_NAVY, PP_ALIGN.CENTER)

# =============================================================
# SLIDE 6: II. Dấu phẩy: Quy tắc gõ văn bản chuẩn (1 Card lớn chữ to)
# =============================================================
print("Slide 6: Dấu phẩy - Quy tắc gõ văn bản chuẩn...")
s6 = prs.slides[5]
clean_slide_content_shapes(s6)
setup_header(s6, "II. Dấu phẩy: Quy tắc gõ văn bản chuẩn và tránh lỗi")

card6 = create_card(s6, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s6_items = [
    "• Quy tắc viết dính sát (Tuyệt đối không cách chữ phía trước):",
    "Dấu phẩy phải được đặt liền sát ký tự cuối cùng của từ đứng trước.",
    "  - ✅ Đúng: \"sách, vở, bút thước\"",
    "  - ❌ Sai: \"sách , vở , bút thước\" (Lỗi gõ thừa dấu cách trước dấu phẩy)",
    "• Quy tắc khoảng trắng (Bắt buộc cách 1 dấu cách phía sau):",
    "Sau dấu phẩy, bắt buộc phải có đúng một khoảng trắng trước khi viết từ tiếp theo.",
    "  - ✅ Đúng: \"Chiều nay, chúng tôi học bài.\"",
    "  - ❌ Sai: \"Chiều nay,chúng tôi học bài.\" (Lỗi dính chữ sau dấu phẩy)",
    "• Quy tắc viết hoa sau dấu phẩy:",
    "Chữ cái đứng ngay sau dấu phẩy KHÔNG viết hoa (trừ danh từ riêng như tên người, tên địa danh)."
]
populate_card_content(card6, "3 Quy tắc soạn thảo văn bản bắt buộc phải nhớ", s6_items, title_size=21, item_size=18, line_spacing=6)

# =============================================================
# SLIDE 7: II. Dấu phẩy: Công dụng và ý nghĩa thực tế (1 Card lớn chữ to)
# =============================================================
print("Slide 7: Dấu phẩy - Công dụng và ý nghĩa thực tế...")
s7 = prs.slides[6]
clean_slide_content_shapes(s7)
setup_header(s7, "II. Dấu phẩy: Công dụng và ý nghĩa trong câu")

card7 = create_card(s7, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s7_items = [
    "• Tạo quãng nghỉ ngắn khi đọc và nói:",
    "Tạo nhịp nghỉ bằng 1/2 nhịp dấu chấm, giúp người nói lấy hơi tự nhiên và người nghe nắm bắt mạch câu dễ dàng.",
    "• Phân định rõ ràng các thành phần trong câu:",
    "Giúp câu văn dài không bị rối, làm nổi bật trạng ngữ, bổ ngữ và các vế câu quan trọng.",
    "• Tránh hiểu lầm nghĩa hoặc gây hiểu sai tai hại:",
    "Đặt dấu phẩy đúng chỗ sẽ phân định ranh giới ngữ nghĩa, loại bỏ các cách hiểu mơ hồ hoặc sai lệch.",
    "• Giúp liệt kê sự vật, hành động mạch lạc:",
    "Ngăn cách các từ ngữ đồng chức, giúp văn bản trình bày rõ ràng, có tính thẩm mỹ và dễ theo dõi."
]
populate_card_content(card7, "4 Công dụng cốt lõi của Dấu Phẩy", s7_items, title_size=21, item_size=19, line_spacing=7)

# =============================================================
# SLIDE 8: II. Dấu phẩy: 8 Ví dụ minh họa thực tế
# =============================================================
print("Slide 8: Dấu phẩy - 8 Ví dụ minh họa thực tế...")
s8 = prs.slides[7]
clean_slide_content_shapes(s8)
setup_header(s8, "II. Dấu phẩy: 8 Ví dụ minh họa thực tế")

card8 = create_card(s8, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s8_items = [
    "Ví dụ 1 (Liệt kê đồ vật): \"Bàn, ghế, tủ và giường đều được làm bằng gỗ.\"",
    "Ví dụ 2 (Tách trạng ngữ thời gian): \"Chiều nay, chúng tôi có buổi học nhóm ở thư viện.\"",
    "Ví dụ 3 (Nối hai vế câu ghép): \"Trời bắt đầu đổ mưa, đường phố trở nên vắng vẻ.\"",
    "Ví dụ 4 (Tách lời gọi đáp, hô ngữ): \"Lan ơi, hãy mở cửa giúp mình với.\"",
    "Ví dụ 5 (Tách thành phần chú thích): \"Nam, lớp trưởng lớp tôi, luôn đạt danh hiệu sinh viên giỏi.\"",
    "Ví dụ 6 (Tách trạng ngữ nơi chốn): \"Trên sân trường rợp bóng cây, các bạn sinh viên đang thảo luận bài.\"",
    "Ví dụ 7 (Liệt kê chuỗi hành động): \"Bạn ấy đọc tài liệu, ghi chép cẩn thận và tích cực phát biểu.\"",
    "Ví dụ 8 (Vế câu nhượng bộ): \"Dù thời gian làm bài khá gấp, cả nhóm vẫn hoàn thành báo cáo đúng hạn.\""
]
populate_card_content(card8, "Hệ thống 8 ví dụ minh họa kèm tác dụng cụ thể", s8_items, title_size=20, item_size=17, line_spacing=6)

# =============================================================
# SLIDE 9: III. Dấu chấm: Vị trí & Quy tắc viết (Card + Ảnh 3D)
# =============================================================
print("Slide 9: Dấu chấm - Vị trí và quy tắc viết...")
s9 = prs.slides[8]
clean_slide_content_shapes(s9)
setup_header(s9, "III. Dấu chấm: Vị trí và quy tắc viết chuẩn")

card9 = create_card(s9, Inches(1.3), Inches(2.80), Inches(9.6), Inches(5.75))
s9_items = [
    "• Luôn đặt ở vị trí cuối cùng của câu kể (câu trần thuật):",
    "Đánh dấu câu đã hoàn chỉnh cả về ngữ pháp và ý nghĩa thông báo.",
    "• Dùng trong câu cầu khiến nhẹ:",
    "Có thể dùng dấu chấm ở cuối câu khuyên nhủ, dặn dò để thể hiện thái độ lịch thiệp.",
    "• Quy tắc viết liền chữ trước:",
    "Viết dính sát ký tự cuối cùng của từ đứng trước. Tuyệt đối không để khoảng trắng trước dấu chấm.",
    "• Quy tắc dấu cách phía sau:",
    "Bắt buộc cách đúng 1 khoảng trắng sau dấu chấm trước khi viết câu mới.",
    "• Bắt buộc viết hoa chữ cái đầu câu tiếp theo:",
    "Chữ cái đầu tiên của từ đứng sau dấu chấm kết thúc câu bắt buộc phải viết hoa."
]
populate_card_content(card9, "Vị trí & Cách đặt Dấu Chấm trong câu", s9_items, title_size=21, item_size=18, line_spacing=6)

card9_img = create_card(s9, Inches(11.3), Inches(2.80), Inches(5.3), Inches(5.75))
s9.shapes.add_picture(r'd:\folder\tools\other\canva\assets\dau_cham.jpg', Inches(11.45), Inches(2.95), Inches(5.0), Inches(5.0))
tb9_lbl = s9.shapes.add_textbox(Inches(11.3), Inches(8.0), Inches(5.3), Inches(0.45))
clear_and_set_text(tb9_lbl.text_frame, "DẤU CHẤM (.)", 'Canva Sans Bold', 17, True, False, COLOR_NAVY, PP_ALIGN.CENTER)

# =============================================================
# SLIDE 10: III. Dấu chấm: Công dụng & Quy ước thường gặp
# =============================================================
print("Slide 10: Dấu chấm - Công dụng và quy ước thường gặp...")
s10 = prs.slides[9]
clean_slide_content_shapes(s10)
setup_header(s10, "III. Dấu chấm: Công dụng và các quy ước thường gặp")

card10 = create_card(s10, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s10_items = [
    "• Đánh dấu kết thúc trọn vẹn một câu kể:",
    "Báo hiệu người viết đã kết thúc trọn vẹn một ý thông tin trước khi chuyển sang câu tiếp theo.",
    "• Quy ước trong từ viết tắt học vị, chức danh:",
    "Dùng sau các từ viết tắt phổ biến: ThS. (Thạc sĩ), TS. (Tiến sĩ), PGS., GS.",
    "• Quy ước phân cách số tiền, số lượng lớn:",
    "Ngăn cách từng cụm 3 chữ số theo quy chuẩn kế toán Việt Nam: 1.000.000 VNĐ, 500.000 đồng.",
    "• Quy ước ghi mốc ngày tháng năm:",
    "Dùng phân định ngày, tháng, năm trong văn bản: 21.09.2026.",
    "• Quy tắc khi kết hợp với dấu ngoặc đơn ():",
    "  - Nếu trong ngoặc là chú thích ngắn: Dấu chấm nằm NGOÀI ngoặc: Bạn cần nộp bài đúng hạn (trước 22h00).",
    "  - Nếu trong ngoặc là câu độc lập hoàn chỉnh: Dấu chấm nằm TRONG ngoặc: (Tất cả tài liệu đã được gửi qua email.)"
]
populate_card_content(card10, "Công dụng chính & 4 Quy ước mở rộng cần nhớ", s10_items, title_size=21, item_size=17, line_spacing=5)

# =============================================================
# SLIDE 11: III. Dấu chấm: 8 Ví dụ minh họa thực tế
# =============================================================
print("Slide 11: Dấu chấm - 8 Ví dụ minh họa thực tế...")
s11 = prs.slides[10]
clean_slide_content_shapes(s11)
setup_header(s11, "III. Dấu chấm: 8 Ví dụ minh họa thực tế")

card11 = create_card(s11, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s11_items = [
    "Ví dụ 1 (Kết thúc câu kể): \"Hôm nay nhóm chúng ta cùng tìm hiểu về quy tắc dùng dấu câu tiếng Việt.\"",
    "Ví dụ 2 (Dặn dò nhẹ nhàng): \"Bạn nhớ kiểm tra lại nội dung bài tập trước khi nộp nhé.\"",
    "Ví dụ 3 (Viết tắt học vị): \"Giảng viên hướng dẫn môn học là ThS. Nguyễn Văn A.\"",
    "Ví dụ 4 (Phân cách số tiền lớn): \"Tổng kinh phí thực hiện đề tài là 1.500.000 đồng.\"",
    "Ví dụ 5 (Ghi mốc ngày tháng năm): \"Thời hạn cuối nộp bài báo cáo là ngày 23.09.2026.\"",
    "Ví dụ 6 (Dấu chấm ngoài ngoặc): \"Sinh viên cần nộp báo cáo đúng hạn (trước 22h00 tối nay).\"",
    "Ví dụ 7 (Dấu chấm trong ngoặc): \"(Mọi thắc mắc về đề tài xin liên hệ trực tiếp với nhóm trưởng.)\"",
    "Ví dụ 8 (Ngắt câu độc lập): \"Mặt trời đã lên cao. Tiếng chim hót ríu rít khắp sân trường.\""
]
populate_card_content(card11, "Hệ thống 8 ví dụ minh họa kèm trường hợp cụ thể", s11_items, title_size=20, item_size=17, line_spacing=6)

# =============================================================
# SLIDE 12: IV. Dấu hỏi: Vị trí & Quy tắc viết (Card + Ảnh 3D)
# =============================================================
print("Slide 12: Dấu hỏi - Vị trí và quy tắc viết (Nội dung của bạn)...")
s12 = prs.slides[11]
clean_slide_content_shapes(s12)
setup_header(s12, "IV. Dấu hỏi: Vị trí và quy tắc viết chuẩn")

card12 = create_card(s12, Inches(1.3), Inches(2.80), Inches(9.6), Inches(5.75))
s12_items = [
    "• Vị trí của dấu hỏi trong câu:",
    "Dấu hỏi (dấu hỏi chấm ?) luôn đứng ở cuối câu hỏi hoặc cuối một vế câu dùng để nghi vấn.",
    "• Quy tắc khoảng trắng khi soạn thảo:",
    "Khi viết, dấu hỏi được đặt ngay sát chữ cái cuối cùng của từ đứng trước (không khoảng trắng) và phía sau có một khoảng trắng trước khi bắt đầu câu mới.",
    "  - ✅ Đúng: \"Bạn đã làm bài tập chưa?\"",
    "  - ❌ Sai: \"Bạn đã làm bài tập chưa ?\" (Lỗi gõ thừa dấu cách)",
    "• Quy tắc viết hoa sau dấu hỏi:",
    "Chữ cái đầu tiên của câu tiếp theo sau dấu hỏi bắt buộc phải viết hoa theo đúng chuẩn chính tả tiếng Việt."
]
populate_card_content(card12, "Vị trí & Quy tắc soạn thảo chuẩn của Dấu Hỏi", s12_items, title_size=21, item_size=18, line_spacing=6)

card12_img = create_card(s12, Inches(11.3), Inches(2.80), Inches(5.3), Inches(5.75))
s12.shapes.add_picture(r'd:\folder\tools\other\canva\assets\dau_hoi.jpg', Inches(11.45), Inches(2.95), Inches(5.0), Inches(5.0))
tb12_lbl = s12.shapes.add_textbox(Inches(11.3), Inches(8.0), Inches(5.3), Inches(0.45))
clear_and_set_text(tb12_lbl.text_frame, "DẤU HỎI (?)", 'Canva Sans Bold', 17, True, False, COLOR_NAVY, PP_ALIGN.CENTER)

# =============================================================
# SLIDE 13: IV. Dấu hỏi: Vai trò & Tác dụng (Nội dung của bạn)
# =============================================================
print("Slide 13: Dấu hỏi - Vai trò và tác dụng (Nội dung của bạn)...")
s13 = prs.slides[12]
clean_slide_content_shapes(s13)
setup_header(s13, "IV. Dấu hỏi: Vai trò và tác dụng thực tế")

card13 = create_card(s13, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s13_items = [
    "• Báo hiệu câu hỏi (nghi vấn):",
    "Giúp người đọc nhận biết câu đó dùng để hỏi, tìm kiếm thông tin hoặc yêu cầu giải đáp.",
    "• Tạo ngữ điệu khi nói và đọc:",
    "Trong văn nói, dấu hỏi biểu thị việc lên giọng ở cuối câu để người nghe dễ dàng nhận biết câu hỏi.",
    "• Thể hiện các mục đích giao tiếp khác (câu hỏi tu từ):",
    "  - Bày tỏ sự băn khoăn, nghi ngờ: Tự hỏi chính mình hoặc hoài nghi về một sự việc.",
    "  - Cầu khiến, nhờ vả, bộc lộ cảm xúc: Mượn hình thức câu hỏi để bộc lộ sự ngạc nhiên, mỉa mai, nhờ giúp đỡ hoặc khẳng định/phủ định một điều gì đó mà không cần người khác trả lời."
]
populate_card_content(card13, "Vai trò và tác dụng đa dạng của Dấu Hỏi", s13_items, title_size=21, item_size=19, line_spacing=7)

# =============================================================
# SLIDE 14: IV. Dấu hỏi: 8 Ví dụ minh họa thực tế (5 Ví dụ của bạn + 3 mở rộng)
# =============================================================
print("Slide 14: Dấu hỏi - 8 Ví dụ minh họa thực tế (Đầy đủ ví dụ của bạn)...")
s14 = prs.slides[13]
clean_slide_content_shapes(s14)
setup_header(s14, "IV. Dấu hỏi: 8 Ví dụ minh họa thực tế")

card14 = create_card(s14, Inches(1.3), Inches(2.80), Inches(15.3), Inches(5.75))
s14_items = [
    "Ví dụ 1 (Hỏi thông tin thời gian): \"Mấy giờ rồi bạn ơi?\"",
    "Ví dụ 2 (Câu hỏi lựa chọn): \"Hôm nay chúng ta ăn cơm hay ăn phở?\"",
    "Ví dụ 3 (Băn khoăn, tự hỏi mình): \"Không biết liệu mình làm như thế này đã đúng chưa?\"",
    "Ví dụ 4 (Cầu khiến, nhờ vả lịch sự): \"Bạn có thể lấy giúp mình quyển sách trên kệ được không?\"",
    "Ví dụ 5 (Cảm xúc ngạc nhiên, trầm trồ): \"Sao cảnh vật quê hương lại đẹp đến thế này?\"",
    "Ví dụ 6 (Câu hỏi tu từ mang ý khẳng định): \"Có ai lại không mong muốn nhóm mình đạt kết quả xuất sắc?\"",
    "Ví dụ 7 (Biểu thị thông tin chưa xác minh): \"Nhà văn sinh năm 1902 (?) và mất năm 1963 tại Hà Nội.\"",
    "Ví dụ 8 (Hỏi lễ phép với thầy cô): \"Thưa thầy, nhóm em có thể xin thêm tài liệu tham khảo được không ạ?\""
]
populate_card_content(card14, "Hệ thống 8 ví dụ minh họa sinh động", s14_items, title_size=20, item_size=17, line_spacing=6)

prs.save(dst_path)
print(f"\nSUCCESS: Successfully built {len(prs.slides)} slides with large text and 3D illustrations in {dst_path}")
